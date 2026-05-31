"""Word 报告生成：使用动态模板样式

格式由 references/output-Demo.docx 定义，运行时动态读取。
包含：风险四象限、因果关系、趋势标记、依赖关系图、整体风险评分
"""

import logging
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn as QN

from .analyzer import PersonAnalysis, RiskItem, DependencyItem
from .report_style import get_report_style, ReportStyle

logger = logging.getLogger(__name__)


def generate_report(
    analyses: list[PersonAnalysis],
    output_path: str | Path,
    chat_name: str = "汇报",
    dependencies: list[DependencyItem] | None = None,
    overall_risk_score: int = 0,
    overall_summary: str = "",
    key_concerns: list[str] | None = None,
    reference_path: str | Path | None = None,
) -> str:
    """生成 Word 报告（格式从参考模板动态读取）"""
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    # 加载样式
    style = get_report_style(reference_path)

    doc = Document()

    # --- 页面设置 ---
    _setup_page(doc, style)

    # --- 全局字体 ---
    _set_chinese_font(doc, style)

    # === 标题 ===
    title = doc.add_heading(f"{chat_name} 差异分析报告", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # === 副标题 ===
    now = datetime.now().strftime("%Y-%m-%d %H:%M")
    subtitle = doc.add_paragraph(f"生成时间：{now} | 分析周期：最近2次双周汇报")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(subtitle.runs[0], size=Pt(10), color=RGBColor(128, 128, 128))

    doc.add_paragraph()

    # === 一、整体摘要 ===
    doc.add_heading("一、整体摘要", level=1)

    total_people = len(analyses)
    total_diffs = sum(len(a.diffs) for a in analyses)
    total_risks = sum(len(a.risks) for a in analyses)

    # 风险四象限统计
    blocking = sum(1 for a in analyses for r in a.risks if r.level == "阻塞")
    delayed = sum(1 for a in analyses for r in a.risks if r.level == "延期")
    resource = sum(1 for a in analyses for r in a.risks if r.level == "资源")
    dependency = sum(1 for a in analyses for r in a.risks if r.level == "依赖")

    summary_text = (
        f"本期共 {total_people} 人汇报，"
        f"{total_diffs} 项关键差异，"
        f"{total_risks} 项风险点需关注"
    )
    doc.add_paragraph(summary_text)

    # 风险分布
    if total_risks > 0:
        p = doc.add_paragraph()
        p.add_run("风险分布：").bold = True
        p.add_run(f" 阻塞×{blocking}  延期×{delayed}  资源×{resource}  依赖×{dependency}")

    # 整体风险评分
    if overall_risk_score > 0:
        p = doc.add_paragraph()
        p.add_run(f"整体风险评分：{overall_risk_score}/5  ")
        if overall_risk_score >= 4:
            run = p.add_run("⚠ 高风险")
            run.font.color.rgb = style.get_risk_color("阻塞")
        elif overall_risk_score >= 3:
            run = p.add_run("⚡ 中等风险")
            run.font.color.rgb = style.get_risk_color("延期")
        else:
            run = p.add_run("✓ 低风险")
            run.font.color.rgb = style.get_diff_color("新增")

    if overall_summary:
        doc.add_paragraph(overall_summary)

    # 重点关注（key_concerns 已是 "问题——建议" 格式）
    if key_concerns:
        doc.add_heading("重点关注事项", level=2)
        for concern in key_concerns:
            p = doc.add_paragraph(style="List Bullet")
            run = p.add_run("⚠ ")
            run.font.color.rgb = style.get_risk_color("阻塞")
            p.add_run(concern)

    # === 二、依赖关系图 ===
    if dependencies:
        doc.add_heading("二、跨人依赖关系", level=1)
        _add_dependency_table(doc, style, dependencies)
        doc.add_paragraph()

    # === 三、各人汇报分析 ===
    section_num = 3 if dependencies else 2
    doc.add_heading(f"{'一二三四五六七八九十'[section_num-1]}、各人汇报分析", level=1)

    for analysis in analyses:
        _add_person_section(doc, style, analysis)

    # === 四 / 三、风险汇总 ===
    section_num += 1
    doc.add_heading(f"{'一二三四五六七八九十'[section_num-1]}、风险汇总", level=1)

    all_risks: list[tuple[str, RiskItem]] = []
    for analysis in analyses:
        for risk in analysis.risks:
            all_risks.append((analysis.person_name, risk))

    all_risks.sort(key=lambda x: style.risk_order.get(x[1].level, 99))

    if all_risks:
        _add_risk_table(doc, style, all_risks)
    else:
        doc.add_paragraph("本期无风险点。")

    doc.save(str(output_path))
    logger.info("报告已生成: %s", output_path)
    return str(output_path)


def _setup_page(doc: Document, style: ReportStyle):
    """设置页面布局"""
    section = doc.sections[0]
    section.page_width = style.page_width
    section.page_height = style.page_height
    section.left_margin = style.margin_left
    section.right_margin = style.margin_right
    section.top_margin = style.margin_top
    section.bottom_margin = style.margin_bottom


def _add_dependency_table(doc: Document, style: ReportStyle, dependencies: list[DependencyItem]):
    """添加依赖关系表"""
    ts = style.dep_table_style
    table = doc.add_table(rows=1, cols=len(ts.headers), style="Table Grid")

    for i, w in enumerate(ts.column_widths):
        table.columns[i].width = w

    hdr = table.rows[0].cells
    for i, label in enumerate(ts.headers):
        hdr[i].text = label
    _set_header_row_style(table, ts.header_fill)

    for dep in dependencies:
        row = table.add_row().cells
        row[0].text = dep.from_person
        row[1].text = dep.to_person
        row[2].text = dep.description


def _add_risk_table(doc: Document, style: ReportStyle, all_risks: list[tuple[str, RiskItem]]):
    """添加风险汇总表"""
    ts = style.risk_table_style
    table = doc.add_table(rows=1, cols=len(ts.headers), style="Table Grid")

    for i, w in enumerate(ts.column_widths):
        table.columns[i].width = w

    hdr = table.rows[0].cells
    for i, label in enumerate(ts.headers):
        hdr[i].text = label
    _set_header_row_style(table, ts.header_fill)

    for person_name, risk in all_risks:
        row = table.add_row().cells
        row[0].text = risk.level
        # 风险等级着色
        for p in row[0].paragraphs:
            for run in p.runs:
                run.font.color.rgb = style.get_risk_color(risk.level)
                run.bold = True
        row[1].text = person_name
        row[2].text = risk.content
        row[3].text = risk.cause or "-"
        row[4].text = style.get_trend_mark(risk.trend) if risk.trend else "-"
        row[5].text = risk.suggestion or "-"


def _add_person_section(doc: Document, style: ReportStyle, analysis: PersonAnalysis):
    """添加单人汇报分析段落"""
    heading_text = analysis.person_name
    if analysis.is_first_report:
        heading_text += "（首次汇报）"
    doc.add_heading(heading_text, level=2)

    # 摘要
    if analysis.summary:
        p = doc.add_paragraph()
        p.add_run("摘要：").bold = True
        p.add_run(analysis.summary)

    # 关键差异
    if analysis.diffs:
        doc.add_paragraph("关键差异：", style="List Bullet")
        for diff in analysis.diffs:
            p = doc.add_paragraph(style="List Bullet 2")
            run = p.add_run(f"[{diff.category}] ")
            run.bold = True
            run.font.color.rgb = style.get_diff_color(diff.category)
            p.add_run(diff.content)
            if diff.impact:
                p2 = doc.add_paragraph(style="List Bullet 2")
                _set_run_font(p2.add_run(f"    影响：{diff.impact}"), size=Pt(9))

    # 风险点
    if analysis.risks:
        doc.add_paragraph("风险点：", style="List Bullet")
        for risk in analysis.risks:
            # 风险等级 + 趋势 + 内容（同一行）
            p = doc.add_paragraph(style="List Bullet 2")
            run = p.add_run(f"[{risk.level}] ")
            run.bold = True
            run.font.color.rgb = style.get_risk_color(risk.level)

            if risk.trend:
                trend_text = style.get_trend_mark(risk.trend)
                _set_run_font(p.add_run(f" {trend_text} "), size=Pt(9))

            p.add_run(risk.content)

            # 原因
            if risk.cause:
                p2 = doc.add_paragraph(style="List Bullet 2")
                _set_run_font(p2.add_run(f"    原因：{risk.cause}"), size=Pt(9))

            # 缓解建议
            if risk.suggestion:
                p3 = doc.add_paragraph(style="List Bullet 2")
                _set_run_font(p3.add_run(f"    建议：{risk.suggestion}"), size=Pt(9))

            # 依赖对象
            if risk.depends_on:
                p4 = doc.add_paragraph(style="List Bullet 2")
                _set_run_font(p4.add_run(f"    依赖方：{', '.join(risk.depends_on)}"), size=Pt(9))


def _set_chinese_font(doc: Document, style: ReportStyle):
    """设置文档所有样式中文字体（Normal + Heading 1/2/Title）"""
    east_asian_font = style.font_name
    east_asian_font_heading = style.font_name_east_asia

    for style_name in ["Normal", "Heading 1", "Heading 2", "Title"]:
        if style_name not in doc.styles:
            continue
        s = doc.styles[style_name]
        font = s.font
        font.name = east_asian_font
        rPr = s.element.get_or_add_rPr()
        rFonts = rPr.find(QN("w:rFonts"))
        if rFonts is None:
            rFonts = s.element.makeelement(QN("w:rFonts"), {})
            rPr.append(rFonts)
        rFonts.set(QN("w:eastAsia"), east_asian_font_heading)


def _set_run_font(run, size: Pt | None = None, color: RGBColor | None = None,
                  bold: bool | None = None):
    """便捷设置 run 字体属性"""
    if size is not None:
        run.font.size = size
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.font.bold = bold


def _set_header_row_style(table, header_fill: str = "D9D9D9"):
    """设置表格表头行加粗 + 灰底"""
    for cell in table.rows[0].cells:
        for p in cell.paragraphs:
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(10)
        # 灰色底色
        shading = cell._element.get_or_add_tcPr()
        shd = shading.makeelement(QN("w:shd"), {
            QN("w:val"): "clear",
            QN("w:color"): "auto",
            QN("w:fill"): header_fill,
        })
        shading.append(shd)
